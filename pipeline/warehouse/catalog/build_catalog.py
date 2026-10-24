"""Convert the 数据采集清单 workbook into the tracked indicators.csv catalog.

Usage:
    python3 -m pipeline.warehouse.catalog.build_catalog path/to/清单.xlsx

The workbook stays the human-maintained source of truth. Re-run this command
whenever the list changes; the CSV diff then shows exactly which indicators
were added, removed or redefined.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

import openpyxl

CATALOG_DIR = Path(__file__).resolve().parent
INDICATORS_CSV = CATALOG_DIR / "indicators.csv"

SHEET = "指标清单"
HEADER = [
    "指标ID",
    "模块",
    "具体数据序列",
    "定义与单位",
    "原始发布频率",
    "周报更新建议",
    "免费程度",
    "数据来源与获取方式",
    "获取限制与口径提示",
    "来源入口",
    "补充入口",
    "本轮核查结论",
    "证据定位／核查边界",
]
FIELDS = [
    "indicator_id",
    "module",
    "name",
    "definition",
    "unit",
    "source_freq",
    "weekly_rule",
    "cost",
    "source_desc",
    "limits",
    "source_url",
    "extra_url",
    "check_status",
    "evidence",
    "catalog_version",
]
_UNIT_RE = re.compile(r"单位[:：]\s*([^。；;]+)")
_DATE_RE = re.compile(r"核查日期[:：]\s*(\d{4}-\d{2}-\d{2})")


def parse_unit(definition: str) -> str:
    match = _UNIT_RE.search(definition or "")
    return match.group(1).strip() if match else ""


def read_workbook(path: Path) -> list[dict[str, str]]:
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    header_row = None
    version = ""
    for row in range(1, 30):
        values = [ws.cell(row, c).value for c in range(1, len(HEADER) + 1)]
        text = " ".join(str(v) for v in values if v)
        if not version and (m := _DATE_RE.search(text)):
            version = m.group(1)
        if values == HEADER:
            header_row = row
            break
    if header_row is None:
        raise ValueError(f"header row {HEADER[:3]}... not found in sheet {SHEET}")

    records: list[dict[str, str]] = []
    seen: set[str] = set()
    for row in range(header_row + 1, ws.max_row + 1):
        indicator_id = ws.cell(row, 1).value
        if not indicator_id:
            continue
        indicator_id = str(indicator_id).strip()
        if indicator_id in seen:
            raise ValueError(f"duplicate indicator id {indicator_id} at row {row}")
        seen.add(indicator_id)
        cells = [ws.cell(row, c) for c in range(1, len(HEADER) + 1)]
        text = ["" if c.value is None else str(c.value).strip() for c in cells]
        source_link = cells[9].hyperlink.target if cells[9].hyperlink else ""
        extra_link = cells[10].hyperlink.target if cells[10].hyperlink else ""
        records.append(
            {
                "indicator_id": indicator_id,
                "module": text[1],
                "name": text[2],
                "definition": text[3],
                "unit": parse_unit(text[3]),
                "source_freq": text[4],
                "weekly_rule": text[5],
                "cost": text[6],
                "source_desc": text[7],
                "limits": text[8],
                "source_url": source_link,
                "extra_url": extra_link,
                "check_status": text[11],
                "evidence": text[12],
                "catalog_version": version,
            }
        )
    return records


def write_catalog(records: list[dict[str, str]], path: Path = INDICATORS_CSV) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(records)


def load_catalog(path: Path = INDICATORS_CSV) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=INDICATORS_CSV)
    args = parser.parse_args(argv)
    records = read_workbook(args.workbook)
    write_catalog(records, args.output)
    print(f"wrote {len(records)} indicators to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
